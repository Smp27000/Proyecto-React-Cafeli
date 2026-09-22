from typing import Optional
from datetime import datetime, date
from io import BytesIO
from decimal import Decimal
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import cast, Date, and_

from app.database import get_db
from app.models.venta import Venta, DetalleVenta
from app.models.usuario import Usuario
from app.auth import require_admin_or_empleado
from app.config import settings

router = APIRouter(prefix="/reportes", tags=["Reportes"])


def _filtrar_ventas_fecha(db: Session, fecha_str: str):
    try:
        fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(400, "Formato de fecha inválido (YYYY-MM-DD)")
    return (
        db.query(Venta)
        .filter(
            and_(
                cast(Venta.fecha_venta, Date) == fecha,
                Venta.estado != "Anulada"
            )
        )
        .order_by(Venta.fecha_venta.desc())
        .all()
    ), fecha


@router.get("/ventas-diario/resumen", summary="Resumen de ventas del día")
def reporte_ventas_resumen(
    fecha: str = Query(..., description="Fecha del reporte (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
    _: Usuario = Depends(require_admin_or_empleado)
):
    ventas, f = _filtrar_ventas_fecha(db, fecha)
    total_ventas = len(ventas)
    total_facturado = Decimal("0.00")
    total_subtotal = Decimal("0.00")
    total_impuestos = Decimal("0.00")
    total_descuentos = Decimal("0.00")
    total_items = 0
    clientes_distintos = set()
    metodo_pago_counts = {}

    for v in ventas:
        total_facturado += v.total
        total_subtotal += v.subtotal
        total_impuestos += v.impuestos
        total_descuentos += v.descuento
        clientes_distintos.add(v.cliente_id)
        mp = v.metodo_pago or "Otro"
        metodo_pago_counts[mp] = metodo_pago_counts.get(mp, 0) + 1
        for d in v.detalles:
            total_items += d.cantidad

    return {
        "success": True,
        "fecha": f.isoformat(),
        "total_ventas": total_ventas,
        "total_clientes": len(clientes_distintos),
        "total_items_vendidos": total_items,
        "subtotal": total_subtotal,
        "impuestos": total_impuestos,
        "descuentos": total_descuentos,
        "total_facturado": total_facturado,
        "metodos_pago": metodo_pago_counts,
        "ventas_count": total_ventas
    }


@router.get("/ventas-diario/pdf", summary="Reporte diario de ventas en PDF")
def reporte_ventas_pdf(
    fecha: str = Query(..., description="Fecha (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
    _: Usuario = Depends(require_admin_or_empleado)
):
    try:
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib import colors
    except ImportError:
        raise HTTPException(500, "reportlab no está instalado")

    ventas, f = _filtrar_ventas_fecha(db, fecha)
    total_facturado = sum((v.total for v in ventas), Decimal("0.00"))

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, leftMargin=0.5 * inch, rightMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    elements = []

    title_style = ParagraphStyle("TitleCustom", parent=styles["Title"], fontSize=20, textColor=colors.HexColor("#6F4E37"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, textColor=colors.HexColor("#3b2a22"))

    elements.append(Paragraph(settings.EMPRESA_NOMBRE, title_style))
    elements.append(Paragraph(f"<b>REPORTE DIARIO DE VENTAS</b>", h2))
    elements.append(Paragraph(f"Fecha: {f.strftime('%A, %d de %B de %Y')}", styles["Normal"]))
    elements.append(Paragraph(f"Empresa: {settings.EMPRESA_NOMBRE} | NIT: {settings.EMPRESA_NIT}", styles["Normal"]))
    elements.append(Spacer(1, 0.25 * inch))

    # Resumen
    resumen_data = [
        ["Total Ventas:", str(len(ventas))],
        ["Total Facturado:", f"${total_facturado:,.2f}"],
        ["Generado:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ]
    rt = Table(resumen_data, colWidths=[2.2 * inch, 2.2 * inch])
    rt.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#faf7f2")),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#6F4E37")),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(rt)
    elements.append(Spacer(1, 0.3 * inch))

    # Detalle
    elements.append(Paragraph("<b>Detalle de Ventas</b>", h2))
    headers = ["N° Venta", "Cliente", "Método Pago", "Hora", "Estado", "Total"]
    rows = [headers]
    for v in ventas:
        cliente = f"{v.cliente.nombres} {v.cliente.apellidos}" if v.cliente else "N/A"
        rows.append([
            str(v.id),
            cliente[:30],
            v.metodo_pago or "N/A",
            v.fecha_venta.strftime("%H:%M"),
            v.estado,
            f"${v.total:,.2f}"
        ])
    if not ventas:
        rows.append(["-", "Sin ventas en esta fecha", "-", "-", "-", "$0.00"])

    dt = Table(rows, colWidths=[0.8 * inch, 2.3 * inch, 1.3 * inch, 0.7 * inch, 0.9 * inch, 1.2 * inch], repeatRows=1)
    dt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6F4E37")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (5, 0), (5, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(dt)

    doc.build(elements)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Reporte_Ventas_{f.isoformat()}.pdf"}
    )


@router.get("/ventas-diario/excel", summary="Reporte diario de ventas en Excel")
def reporte_ventas_excel(
    fecha: str = Query(..., description="Fecha (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
    _: Usuario = Depends(require_admin_or_empleado)
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(500, "openpyxl no está instalado")

    ventas, f = _filtrar_ventas_fecha(db, fecha)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Ventas {f.isoformat()}"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="6F4E37", end_color="6F4E37", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Título
    ws.merge_cells("A1:J1")
    ws["A1"] = settings.EMPRESA_NOMBRE
    ws["A1"].font = Font(bold=True, size=16, color="6F4E37")
    ws["A1"].alignment = center
    ws.merge_cells("A2:J2")
    ws["A2"] = f"REPORTE DIARIO DE VENTAS - {f.isoformat()}"
    ws["A2"].font = Font(bold=True, size=13)
    ws["A2"].alignment = center

    # Encabezados
    headers = [
        "N° Venta", "Fecha", "Hora", "Cliente", "Tipo Doc.", "Número Doc.",
        "Método Pago", "Estado", "Subtotal", "Total"
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    total_facturado = Decimal("0.00")
    for i, v in enumerate(ventas, start=5):
        cliente = v.cliente
        total_facturado += v.total
        row_data = [
            v.id,
            v.fecha_venta.strftime("%Y-%m-%d"),
            v.fecha_venta.strftime("%H:%M:%S"),
            f"{cliente.nombres} {cliente.apellidos}" if cliente else "N/A",
            cliente.tipo_documento if cliente else "N/A",
            cliente.numero_documento if cliente else "N/A",
            v.metodo_pago or "N/A",
            v.estado,
            float(v.subtotal),
            float(v.total),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border = border
            cell.alignment = right if col in (9, 10) else (left if col in (4, 8) else center)

    fila_totales = 5 + len(ventas) + 1
    ws.cell(row=fila_totales, column=8, value="TOTALES:").font = Font(bold=True)
    ws.cell(row=fila_totales, column=9, value=float(sum((v.subtotal for v in ventas), Decimal("0.00")))).font = Font(bold=True)
    ws.cell(row=fila_totales, column=10, value=float(total_facturado)).font = Font(bold=True, color="6F4E37")
    ws.cell(row=fila_totales, column=10).alignment = right

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 12
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 14

    # Sheet 2: Detalle por productos
    ws2 = wb.create_sheet("Detalle Items")
    detalle_headers = ["N° Venta", "Producto/Servicio", "Tipo", "Cantidad", "Precio Unit.", "Descuento", "Subtotal"]
    for col, h in enumerate(detalle_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    fila = 2
    for v in ventas:
        for d in v.detalles:
            data = [
                v.id,
                d.descripcion_item,
                d.tipo_item,
                d.cantidad,
                float(d.precio_unitario),
                float(d.descuento_unitario),
                float(d.subtotal),
            ]
            for col, val in enumerate(data, 1):
                c = ws2.cell(row=fila, column=col, value=val)
                c.border = border
            fila += 1

    widths = [12, 40, 14, 10, 14, 14, 14]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[chr(64 + i)].width = w

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Reporte_Ventas_{f.isoformat()}.xlsx"}
    )
